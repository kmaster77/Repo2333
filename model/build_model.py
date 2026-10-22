"""Build model/financial_model.xlsx with LIVE Excel formulas from model/inputs.csv.

    python3 model/build_model.py                 # writes model/financial_model.xlsx
    python3 model/build_model.py --scenario 1    # same, but scenario selector = pessimistic

The workbook mirrors model_core.run() line by line. Sheets Scenarios, Sensitivity and
part of Viability (criterion 4) hold VALUES computed by model_core (a spreadsheet
cannot rerun itself 50 times without data tables); they are marked in grey italics.

For option scenarios (Agent 7): either
  * put values in the "Override" column of the Inputs sheet (column I), or
  * call build_workbook(path, overrides={...}) / model_core.run(table, overrides={...}).
The option hooks (opt_*) at the bottom of the Inputs sheet add capex, revenue, opex
or a per-m3 price/cost delta without changing any formula.
"""
from __future__ import annotations

import argparse
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import model_core as mc  # noqa: E402
import options as opts  # noqa: E402

OUT = os.path.join(HERE, "financial_model.xlsx")

BLUE = Font(color="0000FF")
BLACK = Font(color="000000")
BOLD = Font(bold=True)
GREY = Font(color="7F7F7F", italic=True)
TITLE = Font(bold=True, size=14)
HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="44546A")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="E7E6E6")
CHECK_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")

EUR = '#,##0;[Red]-#,##0'
EUR2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'
NYRS = mc.N_YEARS


def col(t: int) -> str:
    return get_column_letter(3 + t)


# ---------------------------------------------------------------------------
# Registry of yearly rows: key -> (sheet, row). Templates use tokens:
#   {key}      same column            {key@-1} previous column
#   {key@0}    absolute column C (t=0, used for scalars)
#   {key:cum}  $C$r:<this col>r       {t} / {year}  this column's t / year cell
# Names of inputs (e.g. output_steady_state) are workbook defined names.
# ---------------------------------------------------------------------------
REG: dict = {}
TOK = re.compile(r"\{([A-Za-z0-9_]+)(?:@(-1|0))?(?::(cum))?\}")


class YearSheet:
    def __init__(self, ws, title, note=""):
        self.ws = ws
        self.name = ws.title
        self.rows = []
        self.r = 6
        ws["A1"] = title
        ws["A1"].font = TITLE
        ws["A2"] = note
        ws["A2"].font = GREY
        ws["A3"], ws["A4"] = "Year", "t (year index)"
        ws["A3"].font = ws["A4"].font = BOLD
        ws["B3"] = "unit"
        for t in mc.T:
            ws[f"{col(t)}4"] = t                       # structural constant
            ws[f"{col(t)}3"] = f"={mc.BASE_YEAR}+{col(t)}4"
            for rr in (3, 4):
                c = ws[f"{col(t)}{rr}"]
                c.font, c.fill, c.alignment = HDR, HDR_FILL, Alignment(horizontal="center")
        ws["A3"].fill = ws["A4"].fill = ws["B3"].fill = ws["B4"].fill = HDR_FILL
        ws["A3"].font = ws["A4"].font = ws["B3"].font = HDR
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 12
        for t in mc.T:
            ws.column_dimensions[col(t)].width = 13
        ws.freeze_panes = "C5"

    def section(self, text):
        self.ws.cell(self.r, 1, text).font = Font(bold=True, color="44546A", size=11)
        self.r += 1

    def row(self, key, label, unit, tmpl, fmt=EUR, total=False, check=False, cols="all"):
        REG[key] = (self.name, self.r)
        self.rows.append((key, self.r, label, unit, tmpl, fmt, total, check, cols))
        self.r += 1

    def blank(self):
        self.r += 1

    def write(self):
        ws = self.ws
        for key, r, label, unit, tmpl, fmt, total, check, cols in self.rows:
            ws.cell(r, 1, label)
            ws.cell(r, 2, unit)
            if total:
                ws.cell(r, 1).font = BOLD
            for t in mc.T:
                if cols == "c0" and t != 0:
                    continue
                tp = tmpl(t) if callable(tmpl) else tmpl
                if tp is None:
                    continue
                c = ws.cell(r, 3 + t, "=" + expand(tp, t, self.name) if isinstance(tp, str) else tp)
                c.number_format = fmt
                c.font = BOLD if total else BLACK
                if total:
                    c.fill = TOTAL_FILL
                if check:
                    c.fill = CHECK_FILL


def expand(tmpl: str, t: int, sheet: str) -> str:
    def rep(m):
        key, mod, rng = m.group(1), m.group(2), m.group(3)
        if key == "t":
            return f"{col(t)}$4" if sheet != "__" else str(t)
        if key == "year":
            return f"{col(t)}$3"
        sh, row = REG[key]
        pre = "" if sh == sheet else f"{sh}!"
        if rng == "cum":
            return f"{pre}$C${row}:{col(t)}{row}"
        if mod == "0":
            return f"{pre}$C${row}"
        if mod == "-1":
            return f"{pre}{col(t - 1)}{row}"
        return f"{pre}{col(t)}{row}"
    return TOK.sub(rep, tmpl)


def R(key, t, absolute=True):
    sh, row = REG[key]
    return f"{sh}!${col(t)}${row}" if absolute else f"{sh}!{col(t)}{row}"


def RNG(key, t0=0, t1=NYRS):
    sh, row = REG[key]
    return f"{sh}!${col(t0)}${row}:${col(t1)}${row}"


def ops(tmpl, t0_value=0):
    """Template only for operating years; t=0 gets a constant."""
    return lambda t: tmpl if t >= 1 else t0_value


def only_t0(tmpl):
    return lambda t: tmpl if t == 0 else 0


def hdr(ws, r, labels, c0=1):
    for i, lab in enumerate(labels):
        c = ws.cell(r, c0 + i, lab)
        c.font, c.fill = HDR, HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


# ---------------------------------------------------------------------------
def build_workbook(path=OUT, scenario_selector=2, overrides=None, table=None, with_python_sheets=True, option_selector=0):
    REG.clear()
    table = table or mc.load_inputs()
    overrides = overrides or {}
    wb = Workbook()
    names = ["README", "Inputs", "UnitEconomics", "Capex", "Revenue", "Opex", "PnL", "CashFlow",
             "Valuation", "Scenarios", "Sensitivity", "Viability", "Opt_Overview", "Opt_AssetLight"]
    sheets = {n: (wb.active if i == 0 else wb.create_sheet(n)) for i, n in enumerate(names)}
    sheets["README"].title = "README"

    # ======================= Inputs =========================================
    ws = sheets["Inputs"]
    ws["A1"] = "Inputs - all model parameters (from model/inputs.csv) + option hooks"
    ws["A1"].font = TITLE
    ws["A3"] = "ACTIVE SCENARIO (1 = pessimistic, 2 = base, 3 = optimistic)"
    ws["A3"].font = BOLD
    ws["C3"] = scenario_selector
    ws["C3"].font, ws["C3"].fill = Font(color="0000FF", bold=True, size=12), INPUT_FILL
    ws["D3"] = '=CHOOSE(scenario,"Pessimistic","Base","Optimistic")'
    ws["D3"].font = BOLD
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C3")
    wb.defined_names["scenario"] = DefinedName("scenario", attr_text="Inputs!$C$3")
    ws["F3"] = "ACTIVE OPTION (0 = none; see Opt_Overview)"
    ws["F3"].font = BOLD
    ws["I3"] = option_selector
    ws["I3"].font, ws["I3"].fill = Font(color="0000FF", bold=True, size=12), INPUT_FILL
    wb.defined_names["option_sel"] = DefinedName("option_sel", attr_text="Inputs!$I$3")
    opt_params = []
    for o in opts.MODEL_OPTIONS:
        for k_ in o[3]:
            if k_ not in opt_params:
                opt_params.append(k_)
    OPT_R0, OPT_R1 = 7, 7 + len(opt_params) - 1
    OPT_C1 = get_column_letter(2 + len(opts.MODEL_OPTIONS))
    ws["A4"] = ("Column H picks low/base/high per row: pessimistic uses column F, optimistic column G "
                "(1=low, 2=base, 3=high). Column I = OVERRIDE (leave empty = use H). Column O = value of the ACTIVE OPTION "
                "(cell I3, matrix on sheet Opt_Overview). Column J = ACTIVE value used by every formula "
                "(priority: Override > option > scenario; defined name = parameter name).")
    ws["A4"].font = GREY
    hdr(ws, 5, ["parameter (= defined name)", "unit", "low", "base", "high", "pess. col", "opt. col",
                "scenario value", "OVERRIDE", "ACTIVE", "label", "source_ids", "notes", "direction", "option value"])
    r = 6
    rows = list(table.items()) + [(h, {"unit": u, "low": 0.0, "base": 0.0, "high": 0.0, "label": "MODEL HOOK",
                                       "source_ids": "model", "notes": d}) for h, (u, d) in mc.HOOKS.items()]
    for p, rec in rows:
        if p in mc.HOOKS and r == 6 + len(table):
            ws.cell(r, 1, "Option hooks (not in inputs.csv; 0 = no effect)").font = Font(bold=True, color="44546A")
            r += 1
        d = mc.DIRECTION.get(p, "fixed")
        ws.cell(r, 1, p)
        ws.cell(r, 2, rec.get("unit"))
        for j, c in enumerate(mc.COLS):
            cell = ws.cell(r, 3 + j, rec[c])
            cell.font = BLUE
        ws.cell(r, 6, mc.scenario_col(p, "pessimistic")).font = BLUE
        ws.cell(r, 7, mc.scenario_col(p, "optimistic")).font = BLUE
        ws.cell(r, 8, f"=INDEX(C{r}:E{r},1,CHOOSE(scenario,F{r},2,G{r}))")
        ov = ws.cell(r, 9, overrides.get(p))
        ov.font, ov.fill = BLUE, INPUT_FILL
        mtx = f"Opt_Overview!$C${OPT_R0}:${OPT_C1}${OPT_R1}"
        mt = f"MATCH($A{r},Opt_Overview!$A${OPT_R0}:$A${OPT_R1},0)"
        ws.cell(r, 15, f'=IFERROR(IF(option_sel=0,"",IF(LEN(INDEX({mtx},{mt},option_sel))=0,"",INDEX({mtx},{mt},option_sel))),"")')
        a = ws.cell(r, 10, f'=IF(I{r}<>"",I{r},IF(O{r}<>"",O{r},H{r}))')
        a.font = BOLD
        ws.cell(r, 11, rec.get("label"))
        ws.cell(r, 12, rec.get("source_ids"))
        nt = rec.get("notes") or ""
        ws.cell(r, 13, ("'" + nt) if nt.startswith("=") else nt)
        ws.cell(r, 14, d)
        wb.defined_names[p] = DefinedName(p, attr_text=f"Inputs!$J${r}")
        r += 1
    for k in overrides:
        if k not in dict(rows):
            raise KeyError(k)
    for c, w in zip("ABCDEFGHIJKLMN", (44, 12, 11, 11, 11, 8, 8, 12, 11, 12, 11, 22, 60, 9)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "B6"

    # ======================= Capex ==========================================
    cx = YearSheet(sheets["Capex"], "Capex, one-off start-up costs, grant, depreciation, salvage (EUR, real 2026)",
                   "t=0 (2026) = permitting & construction year. All initial spend in t=0.")
    cx.section("Initial investment (t = 0)")
    cx.row("capex_plant", "Plant equipment (batching plant, mixer, bins, silos, lab) - in build year (2026 + delay)", "EUR",
           "IF({t}=opt_start_delay,capex_plant_equipment*1000000,0)")
    cx.row("capex_civil", "Site preparation & civil works - in build year", "EUR", "IF({t}=opt_start_delay,capex_site_prep_civil*1000000,0)")
    cx.row("capex_permits", "Permits, studies, engineering (amortised) - always 2026", "EUR", only_t0("capex_permits_engineering*1000000"))
    cx.row("capex_extra", "Option hook: extra capex - in build year", "EUR", "IF({t}=opt_start_delay,opt_extra_capex_y0,0)")
    cx.row("capex_land", "Land purchase (0 = leased) - in build year", "EUR", "IF({t}=opt_start_delay,land_purchase*1000000,0)")
    cx.row("crush_t", "Option A+B: year index of crushing-line capex (year before first crushing)", "t",
           f"MIN({NYRS},MAX(0,ROUND(opt_crusher_start_year-{mc.BASE_YEAR},0)-1))", fmt="0", cols="c0")
    cx.row("capex_crusher", "Option A+B: crushing-line capex", "EUR", "IF({t}={crush_t@0},opt_crusher_capex,0)")
    cx.row("capex_initial", "Total initial capex", "EUR",
           "{capex_plant}+{capex_civil}+{capex_permits}+{capex_extra}+{capex_land}+{capex_crusher}", total=True)
    cx.row("physical", "Physical assets (plant + civil) - base for maintenance, sustaining capex, salvage", "EUR",
           f"SUM({RNG('capex_plant')})+SUM({RNG('capex_civil')})", cols="c0")
    cx.blank()
    cx.section("One-off items in t = 0 (expensed / income, not capex)")
    cx.row("startup_oneoff", "Start-up costs expensed: LCA/EPD + pre-opening (build year)", "EUR",
           "IF({t}=opt_start_delay,capex_lca_epd_one_off+capex_preopening_costs,0)")
    cx.row("grant", "Grant DEI+/VEKI (taxable income in build year)", "EUR", "IF({t}=opt_start_delay,grant_dei_veki_first_plant*1000000,0)")
    cx.row("mia", "MIA extra tax deduction (gives the stated net value at top rate); plant in build year, crusher in its capex year", "EUR",
           "IF({t}=opt_start_delay,mia_eligible_share_of_capex/100*{physical@0}*mia_net_value_pct_eligible_capex/tax_rate_vpb_high_bracket,0)"
           "+IF({t}={crush_t@0},opt_crusher_mia_share_pct/100*opt_crusher_capex*mia_net_value_pct_eligible_capex/tax_rate_vpb_high_bracket,0)")
    cx.blank()
    cx.section("Sustaining capex and depreciation")
    S0 = "(1+opt_start_delay)"      # first operating year index
    cx.row("sust", "Sustaining capex (% of physical assets per operating year)", "EUR",
           f"IF({{t}}>={S0},sustaining_capex_pct_of_capex/100*{{physical@0}},0)")
    cx.row("dep_plant", "Depreciation plant equipment (+ hook capex), from first operating year", "EUR",
           f"IF(AND({{t}}>={S0},{{t}}<{S0}+lifetime_plant_equipment),(SUM({RNG('capex_plant')})+SUM({RNG('capex_extra')}))/lifetime_plant_equipment,0)")
    cx.row("dep_civil", "Depreciation civil works", "EUR",
           f"IF(AND({{t}}>={S0},{{t}}<{S0}+lifetime_civil_works),SUM({RNG('capex_civil')})/lifetime_civil_works,0)")
    cx.row("dep_permits", "Amortisation permits & engineering", "EUR",
           f"IF(AND({{t}}>={S0},{{t}}<{S0}+lifetime_permits_engineering),{{capex_permits@0}}/lifetime_permits_engineering,0)")
    cx.row("dep_crusher", "Depreciation crushing line (option A+B)", "EUR",
           "IF(AND({t}>={crush_t@0}+1,{t}<{crush_t@0}+1+hybrid_crusher_lifetime),opt_crusher_capex/hybrid_crusher_lifetime,0)")
    cx.row("dep_sust", "Depreciation sustaining capex (plant life)", "EUR", "SUM({sust:cum})/lifetime_plant_equipment")
    cx.row("da", "Total depreciation & amortisation", "EUR",
           "{dep_plant}+{dep_civil}+{dep_permits}+{dep_crusher}+{dep_sust}", total=True)
    cx.row("nbv", "Memo: net book value end of year (excl. land)", "EUR",
           "SUM({capex_initial:cum})-SUM({capex_land:cum})+SUM({sust:cum})-SUM({da:cum})")
    cx.row("salvage", "Salvage value in t=10 (= the ONLY terminal value; plant+civil and crushing line)", "EUR",
           f"IF({{t}}={NYRS},salvage_value_yr10_pct_of_plant_civil_capex/100*{{physical@0}}+opt_crusher_salvage_pct/100*opt_crusher_capex,0)")

    # ======================= Revenue ========================================
    rv = YearSheet(sheets["Revenue"], "Volume, prices and revenue",
                   "Plant-average price over all strength classes (Agent 5). Premium only on circular volume.")
    rv.section("Volume")
    rv.row("ramp", "Utilisation of steady-state output (ramp-up)", "%",
           "IF({t}<1+opt_start_delay,0,IF({t}=1+opt_start_delay,ramp_year1_pct_of_steady/100,"
           "IF({t}=2+opt_start_delay,ramp_year2_pct_of_steady/100,ramp_year3plus_pct_of_steady/100)))", fmt=PCT)
    rv.row("vol", "Volume sold", "m3", "output_steady_state*{ramp}", total=True)
    rv.row("vol_circ", "  of which circular mixes (50% coarse RCA)", "m3", "{vol}*circular_share_of_volume/100")
    rv.row("vol_conv", "  of which standard CEM III/B mixes", "m3", "{vol}-{vol_circ}")
    rv.blank()
    rv.section("Cement price path (CEM III/B, delivered)")
    rv.row("cc", "Carbon cost in CEM III/B price (ETS/CBAM)", "EUR/t", lambda t: f"carbon_cost_cem3b_{mc.YEARS[t]}", fmt=EUR2)
    rv.row("slag", "Slag-scarcity uplift from 2031", "EUR/t",
           "IF({year}>=2031,cem3b_slag_scarcity_uplift_2031plus/100*price_cem3b_2026_delivered,0)", fmt=EUR2)
    rv.row("cem_price", "CEM III/B price", "EUR/t",
           "price_cem3b_2026_delivered+cement_carbon_passthrough/100*({cc}-{cc@0})+{slag}", fmt=EUR2, total=True)
    rv.row("mkt_binder", "Market-wide binder cost increase vs 2026 (avg mix)", "EUR/m3",
           "mix_plant_avg_binder_conv/1000*({cem_price}-price_cem3b_2026_delivered)", fmt=EUR2)
    rv.blank()
    rv.section("Prices (delivered, plant average)")
    rv.row("price_market", "Market price, plant average (incl. pass-through of binder cost)", "EUR/m3",
           "price_plant_avg_all_classes*(1+price_real_escalation/100)^{t}"
           "+concrete_price_passthrough_of_market_binder_cost/100*{mkt_binder}+opt_price_delta_per_m3", fmt=EUR2)
    rv.row("premium", "Blended green premium on circular volume", "%",
           "public_share_of_sales/100*green_premium_public/100+(1-public_share_of_sales/100)*green_premium_private/100",
           fmt='0.00%', cols="c0")
    rv.row("price_circ", "Price circular mixes (+ option: circular-only delta from its start year)", "EUR/m3",
           "{price_market}*(1+{premium@0})+IF({year}>=opt_circ_delta_start_year,opt_circ_price_delta_per_m3,0)", fmt=EUR2)
    rv.row("price_conv", "Price standard mixes", "EUR/m3", "{price_market}", fmt=EUR2)
    rv.blank()
    rv.section("Revenue")
    rv.row("rev_circ", "Revenue circular mixes", "EUR", "{vol_circ}*{price_circ}")
    rv.row("rev_conv", "Revenue standard mixes", "EUR", "{vol_conv}*{price_conv}")
    rv.row("rev_extra", "Option hook: extra revenue", "EUR", "IF({t}>=1+opt_start_delay,opt_extra_revenue_per_yr,0)")
    rv.blank()
    rv.section("Option A+B: on-site rubble acceptance and crushing (0 unless opt_crusher_on = 1)")
    rv.row("cflag", "Crushing active (operating year and >= crusher start year)", "-",
           "opt_crusher_on*IF({t}>=1+opt_start_delay,1,0)*IF({year}>=opt_crusher_start_year,1,0)", fmt="0")
    rv.row("intake", "Rubble intake (scales with ramp-up)", "t", "opt_rubble_intake_tpy*{ramp}*{cflag}")
    rv.row("coarse_prod", "Coarse RCA produced", "t", "{intake}*hybrid_yield_coarse/100")
    rv.row("rca_need", "Coarse RCA needed by the plant", "t",
           "{vol_circ}*mix_circ_rca_coarse/1000+{vol_conv}*mix_conv_cem3b_rca_coarse/1000")
    rv.row("rca_own", "Own RCA used in the plant (replaces purchased RCA)", "t", "MIN({coarse_prod},{rca_need})")
    rv.row("rca_surplus", "Surplus coarse RCA sold ex-works", "t", "{coarse_prod}-{rca_own}")
    rv.row("cr_gate", "Gate fee received", "EUR", "{intake}*hybrid_gate_fee_received")
    rv.row("cr_sand", "Crusher sand 0/4 sold", "EUR", "{intake}*hybrid_yield_sand_0_4/100*hybrid_price_crusher_sand")
    rv.row("cr_steel", "Steel scrap sold", "EUR", "{intake}*hybrid_yield_steel/100*hybrid_price_steel_scrap")
    rv.row("cr_surplus", "Surplus RCA sold (at RCA ex-works price)", "EUR", "{rca_surplus}*price_rca_coarse_exworks")
    rv.row("rev_crusher", "Revenue on-site crushing", "EUR", "{cr_gate}+{cr_sand}+{cr_steel}+{cr_surplus}", total=True)
    rv.blank()
    rv.row("revenue", "Total revenue", "EUR", "{rev_circ}+{rev_conv}+{rev_extra}+{rev_crusher}", total=True)
    rv.row("avg_price", "Memo: average realised price", "EUR/m3", lambda t: "IF({vol}>0,({rev_circ}+{rev_conv})/{vol},0)", fmt=EUR2)

    # ======================= Opex ===========================================
    ox = YearSheet(sheets["Opex"], "Operating costs (EUR, real 2026)",
                   "Materials priced per m3 from the recipes; fixed costs from operating year 1.")
    ox.section("Unit inputs")
    ox.row("rca_deliv", "RCA coarse delivered = ex-works + km x rate + handling", "EUR/t",
           "price_rca_coarse_exworks+rca_transport_distance*road_haulage_rate+road_handling_per_t", fmt=EUR2, cols="c0")
    ox.row("rca_qc_m3", "RCA lot testing + handling/pre-wetting (circular only)", "EUR/m3",
           "mix_circ_rca_coarse/1000*(rca_lot_testing+rca_handling_prewetting)", fmt=EUR2, cols="c0")
    ox.row("energy_m3", "Energy (electricity + diesel)", "EUR/m3",
           "electricity_use*price_electricity+diesel_use_plant*price_diesel", fmt=EUR2, cols="c0")
    ox.row("delivery_m3", "Delivery by subcontracted truck mixer", "EUR/m3",
           "delivery_fixed_per_m3+delivery_variable_per_km_oneway*delivery_avg_distance_oneway", fmt=EUR2, cols="c0")
    ox.blank()
    ox.section("Raw materials per m3 - CIRCULAR mix (plant-average binder)")
    ox.row("c_cement", "Cement CEM III/B", "EUR/m3", "mix_plant_avg_binder_circ/1000*{cem_price}", fmt=EUR2)
    ox.row("c_gravel", "Gravel", "EUR/m3", "mix_circ_gravel/1000*price_gravel_delivered", fmt=EUR2)
    ox.row("c_rca", "Recycled aggregate (coarse delivered + fine ex-works)", "EUR/m3",
           "mix_circ_rca_coarse/1000*{rca_deliv@0}+mix_circ_rca_fine/1000*price_rca_fine_exworks", fmt=EUR2)
    ox.row("c_sand", "Sand", "EUR/m3", "mix_circ_sand/1000*price_sand_delivered", fmt=EUR2)
    ox.row("c_admix", "Superplasticiser", "EUR/m3",
           "mix_plant_avg_binder_circ*mix_circ_superplasticiser_pct_binder/100*price_superplasticiser", fmt=EUR2)
    ox.row("c_water", "Water (effective + absorption)", "EUR/m3",
           "(mix_circ_water_effective+mix_circ_water_absorption)/1000*price_water", fmt=EUR2)
    ox.row("mat_circ", "Raw materials circular", "EUR/m3",
           "{c_cement}+{c_gravel}+{c_rca}+{c_sand}+{c_admix}+{c_water}", fmt=EUR2, total=True)
    ox.section("Raw materials per m3 - STANDARD CEM III/B mix")
    ox.row("v_cement", "Cement CEM III/B", "EUR/m3", "mix_plant_avg_binder_conv/1000*{cem_price}", fmt=EUR2)
    ox.row("v_gravel", "Gravel", "EUR/m3", "mix_conv_cem3b_gravel/1000*price_gravel_delivered", fmt=EUR2)
    ox.row("v_rca", "Recycled aggregate", "EUR/m3", "mix_conv_cem3b_rca_coarse/1000*{rca_deliv@0}", fmt=EUR2)
    ox.row("v_sand", "Sand", "EUR/m3", "mix_conv_cem3b_sand/1000*price_sand_delivered", fmt=EUR2)
    ox.row("v_admix", "Superplasticiser", "EUR/m3",
           "mix_plant_avg_binder_conv*mix_conv_cem3b_superplasticiser_pct_binder/100*price_superplasticiser", fmt=EUR2)
    ox.row("v_water", "Water", "EUR/m3",
           "(mix_conv_cem3b_water_effective+mix_conv_cem3b_water_absorption)/1000*price_water", fmt=EUR2)
    ox.row("mat_conv", "Raw materials standard", "EUR/m3",
           "{v_cement}+{v_gravel}+{v_rca}+{v_sand}+{v_admix}+{v_water}", fmt=EUR2, total=True)
    ox.blank()
    ox.section("Indices")
    ox.row("opflag", "Operating year flag", "-", "IF({t}>=1+opt_start_delay,1,0)", fmt="0")
    ox.row("wage_idx", "Real wage index", "x", "(1+real_wage_growth/100)^{t}", fmt="0.0000")
    ox.blank()
    ox.section("Variable costs (EUR/yr)")
    ox.row("o_raw", "Raw materials", "EUR", "{vol_circ}*{mat_circ}+{vol_conv}*{mat_conv}")
    ox.row("o_rcaqc", "RCA testing & handling", "EUR", "{vol_circ}*{rca_qc_m3@0}")
    ox.row("o_energy", "Energy", "EUR", "{vol}*{energy_m3@0}")
    ox.row("o_delivery", "Delivery (truck mixers)", "EUR", "{vol}*{delivery_m3@0}")
    ox.row("o_vardelta", "Option hook: variable cost delta", "EUR", "{vol}*opt_cost_delta_per_m3")
    ox.row("o_circdelta", "Option: circular-only variable cost delta (from its start year)", "EUR",
           "{vol_circ}*IF({year}>=opt_circ_delta_start_year,1,0)*opt_circ_cost_delta_per_m3")
    ox.row("o_cr_rca", "A+B: own RCA replaces purchased RCA (credit)", "EUR", "-{rca_own}*{rca_deliv@0}")
    ox.row("o_cr_cont", "A+B: contaminant disposal", "EUR", "{intake}*hybrid_yield_contaminants/100*hybrid_cost_contaminant_disposal")
    ox.row("o_cr_diesel", "A+B: crusher diesel", "EUR", "{intake}*hybrid_crusher_diesel*price_diesel")
    ox.row("o_cr_wear", "A+B: crusher wear", "EUR", "{intake}*hybrid_crusher_wear")
    ox.row("o_cr_contract", "A+B: contract crushing", "EUR", "{intake}*opt_crusher_contract_rate")
    ox.row("o_baddebt", "Bad debt", "EUR", "{revenue}*bad_debt_pct_revenue/100")
    ox.row("opex_var", "Total variable costs", "EUR",
           "{o_raw}+{o_rcaqc}+{o_energy}+{o_delivery}+{o_vardelta}+{o_circdelta}+{o_cr_rca}+{o_cr_cont}+{o_cr_diesel}"
           "+{o_cr_wear}+{o_cr_contract}+{o_baddebt}", total=True)
    ox.section("Fixed costs (EUR/yr, from year 1)")
    ox.row("o_staffprod", "Production staff", "EUR",
           "{opflag}*staff_production_fte*cost_per_fte_production*{wage_idx}")
    ox.row("o_maint", "Maintenance & wear", "EUR", "{opflag}*maintenance_pct_of_capex/100*{physical@0}")
    ox.row("o_cert", "Certification & testing (KOMO, CSC, NMD)", "EUR", "{opflag}*certification_testing_fixed")
    ox.row("o_lease", "Site lease", "EUR", "{opflag}*site_area*site_lease_rate")
    ox.row("o_staffsga", "SG&A staff", "EUR", "{opflag}*staff_sga_fte*cost_per_fte_sga*{wage_idx}")
    ox.row("o_sgaother", "SG&A non-staff", "EUR", "{opflag}*sga_fixed_non_staff")
    ox.row("o_extrafixed", "Option hook: extra fixed opex", "EUR", "{opflag}*opt_extra_opex_per_yr")
    ox.row("o_cr_staff", "A+B: crushing staff", "EUR", "{cflag}*hybrid_crusher_staff_fte*cost_per_fte_production*{wage_idx}")
    ox.row("o_cr_maint", "A+B: crusher maintenance", "EUR", "{cflag}*hybrid_crusher_maintenance_pct_capex/100*opt_crusher_capex")
    ox.row("o_cr_cert", "A+B: BRL 2506 certification", "EUR", "{cflag}*hybrid_brl2506_certification")
    ox.row("o_cr_site", "A+B: extra site area lease", "EUR", "{cflag}*hybrid_extra_site_area*site_lease_rate")
    ox.row("opex_fixed", "Total fixed costs", "EUR",
           "{o_staffprod}+{o_maint}+{o_cert}+{o_lease}+{o_staffsga}+{o_sgaother}+{o_extrafixed}"
           "+{o_cr_staff}+{o_cr_maint}+{o_cr_cert}+{o_cr_site}", total=True)
    ox.row("o_startup", "One-off start-up costs (t=0)", "EUR", "{startup_oneoff}")
    ox.row("opex_total", "TOTAL OPERATING COSTS", "EUR", "{opex_var}+{opex_fixed}+{o_startup}", total=True)
    ox.section("Option A+B memo")
    ox.row("cr_var", "A+B variable costs (incl. RCA credit)", "EUR", "{o_cr_rca}+{o_cr_cont}+{o_cr_diesel}+{o_cr_wear}+{o_cr_contract}")
    ox.row("cr_fix", "A+B fixed costs", "EUR", "{o_cr_staff}+{o_cr_maint}+{o_cr_cert}+{o_cr_site}")
    ox.row("crusher_net", "A+B net contribution (revenue - costs, incl. RCA credit; before bad debt)", "EUR",
           "{rev_crusher}-{cr_var}-{cr_fix}", total=True)

    # ======================= PnL ============================================
    pl = YearSheet(sheets["PnL"], "Profit & loss (EUR, real 2026)", "Unlevered: no debt, no interest.")
    pl.row("p_rev_circ", "Revenue circular mixes", "EUR", "{rev_circ}")
    pl.row("p_rev_conv", "Revenue standard mixes", "EUR", "{rev_conv}")
    pl.row("p_rev_extra", "Other revenue (option hook)", "EUR", "{rev_extra}")
    pl.row("p_rev_crusher", "On-site crushing revenue (option A+B)", "EUR", "{rev_crusher}")
    pl.row("p_revenue", "REVENUE", "EUR", "{p_rev_circ}+{p_rev_conv}+{p_rev_extra}+{p_rev_crusher}", total=True)
    lines = [("o_raw", "Raw materials"), ("o_rcaqc", "RCA testing & handling"), ("o_energy", "Energy"),
             ("o_delivery", "Delivery"), ("o_vardelta", "Variable cost delta (hook)"),
             ("o_circdelta", "Circular-only cost delta (option)"), ("o_cr_rca", "A+B: own RCA credit"),
             ("o_cr_cont", "A+B: contaminant disposal"), ("o_cr_diesel", "A+B: crusher diesel"), ("o_cr_wear", "A+B: crusher wear"),
             ("o_cr_contract", "A+B: contract crushing"), ("o_baddebt", "Bad debt"),
             ("o_staffprod", "Production staff"), ("o_maint", "Maintenance"), ("o_cert", "Certification & testing"),
             ("o_lease", "Site lease"), ("o_staffsga", "SG&A staff"), ("o_sgaother", "SG&A non-staff"),
             ("o_extrafixed", "Extra fixed opex (hook)"), ("o_cr_staff", "A+B: crushing staff"), ("o_cr_maint", "A+B: crusher maintenance"),
             ("o_cr_cert", "A+B: BRL 2506 certification"), ("o_cr_site", "A+B: extra site lease"), ("o_startup", "One-off start-up costs")]
    for k, lab in lines:
        pl.row("p_" + k, "  - " + lab, "EUR", "-{" + k + "}")
    pl.row("ebitda", "EBITDA", "EUR", "{p_revenue}+" + "+".join("{p_" + k + "}" for k, _ in lines), total=True)
    pl.row("ebitda_margin", "EBITDA margin", "%", "IF({p_revenue}<>0,{ebitda}/{p_revenue},0)", fmt=PCT)
    pl.row("p_da", "  - Depreciation & amortisation", "EUR", "-{da}")
    pl.row("ebit", "EBIT", "EUR", "{ebitda}+{p_da}", total=True)
    pl.row("p_grant", "  + Grant income", "EUR", "{grant}")
    pl.row("pbt", "PROFIT BEFORE TAX", "EUR", "{ebit}+{p_grant}", total=True)
    pl.section("Corporate income tax (Vpb)")
    pl.row("p_mia", "MIA deduction", "EUR", "{mia}")
    pl.row("ti", "Taxable result before loss offset", "EUR", "{pbt}-{p_mia}")
    pl.row("taxable", "Taxable profit after loss offset", "EUR",
           lambda t: "MAX(0,{ti})" if t == 0 else "MAX(0,{ti}-{pool@-1})")
    pl.row("pool", "Tax-loss pool carried forward (end of year)", "EUR",
           lambda t: "MAX(0,-{ti})" if t == 0 else "MAX(0,{pool@-1}-{ti})")
    pl.row("tax", "Corporate income tax", "EUR",
           "MIN({taxable},tax_vpb_low_bracket_limit)*tax_rate_vpb_low_bracket/100"
           "+MAX(0,{taxable}-tax_vpb_low_bracket_limit)*tax_rate_vpb_high_bracket/100")
    pl.row("net_income", "NET INCOME", "EUR", "{pbt}-{tax}", total=True)
    pl.section("Checks (must be 0)")
    pl.row("chk_rev", "Check: revenue lines - Revenue sheet total", "EUR", "{p_revenue}-{revenue}", check=True)
    pl.row("chk_opex", "Check: sum of cost lines + Opex total", "EUR",
           "+".join("{p_" + k + "}" for k, _ in lines) + "+{opex_total}", check=True)
    pl.row("chk_ebitda", "Check: EBITDA - (revenue - total opex)", "EUR", "{ebitda}-({revenue}-{opex_total})", check=True)

    # ======================= CashFlow =======================================
    cf = YearSheet(sheets["CashFlow"], "Free cash flow (unlevered, EUR real 2026)",
                   "Terminal value = salvage only (t=10). Working capital is not released at the end.")
    cf.row("f_ebitda", "EBITDA", "EUR", "{ebitda}")
    cf.row("f_grant", "Grant received", "EUR", "{grant}")
    cf.row("f_tax", "Tax paid (same year)", "EUR", "-{tax}")
    cf.section("Working capital")
    cf.row("wc_rec", "Receivables", "EUR", "{revenue}*working_capital_receivable_days/365")
    cf.row("wc_inv", "Inventory (raw materials)", "EUR", "{o_raw}*working_capital_inventory_days/365")
    cf.row("wc_pay", "Payables (raw materials)", "EUR", "{o_raw}*working_capital_payable_days/365")
    cf.row("wc", "Net working capital", "EUR", "{wc_rec}+{wc_inv}-{wc_pay}", total=True)
    cf.row("f_dwc", "Change in working capital (cash out = negative)", "EUR",
           lambda t: "-{wc}" if t == 0 else "-({wc}-{wc@-1})")
    cf.section("Investment")
    cf.row("f_capex", "Initial capex", "EUR", "-{capex_initial}")
    cf.row("f_sust", "Sustaining capex", "EUR", "-{sust}")
    cf.row("f_salvage", "Salvage value (terminal value)", "EUR", "{salvage}")
    cf.row("fcf", "FREE CASH FLOW", "EUR", "{f_ebitda}+{f_grant}+{f_tax}+{f_dwc}+{f_capex}+{f_sust}+{f_salvage}", total=True)
    cf.row("cum", "Cumulative free cash flow", "EUR", "SUM({fcf:cum})")
    cf.row("df", "Discount factor", "x", "1/(1+discount_rate_real/100)^{t}", fmt="0.0000")
    cf.row("dcf", "Discounted free cash flow", "EUR", "{fcf}*{df}")
    cf.row("cum_dcf", "Cumulative discounted FCF", "EUR", "SUM({dcf:cum})")
    cf.section("Payback helpers (cash flow EXCLUDING salvage)")
    cf.row("fcf_x", "FCF excl. salvage", "EUR", "{fcf}-{f_salvage}")
    cf.row("cum_x", "Cumulative FCF excl. salvage", "EUR", "SUM({fcf_x:cum})")
    cf.row("cum_dx", "Cumulative discounted FCF excl. salvage", "EUR", "SUMPRODUCT({fcf_x:cum},{df:cum})")
    cf.row("pb_h", "Payback crossing (year, interpolated)", "yr",
           lambda t: None if t == 0 else 'IF(AND({cum_x@-1}<0,{cum_x}>=0),{t}-1+(-{cum_x@-1})/({cum_x}-{cum_x@-1}),"")', fmt="0.00")
    cf.row("dpb_h", "Discounted payback crossing (year)", "yr",
           lambda t: None if t == 0 else 'IF(AND({cum_dx@-1}<0,{cum_dx}>=0),{t}-1+(-{cum_dx@-1})/({cum_dx}-{cum_dx@-1}),"")', fmt="0.00")
    cf.section("Reconciliation (must be 0)")
    cf.row("chk_cf", "Check: FCF - (net income + D&A + dWC - capex - sustaining + salvage)", "EUR",
           "{fcf}-({net_income}+{da}+{f_dwc}+{f_capex}+{f_sust}+{f_salvage})", check=True)
    cf.row("chk_cf2", "Check: EBITDA + grant - tax - (net income + D&A)", "EUR",
           "{f_ebitda}+{f_grant}+{f_tax}-({net_income}+{da})", check=True)

    for s in (cx, rv, ox, pl, cf):
        s.write()

    # ======================= Valuation ======================================
    va = sheets["Valuation"]
    va["A1"] = "Valuation & KPIs (active scenario)"
    va["A1"].font = TITLE
    va["A2"] = '="Active scenario: "&Inputs!D3'
    hdr(va, 4, ["KPI", "value", "unit", "how it is calculated"])
    kp = {}

    def kpi(key, label, formula, unit, how, fmt=EUR):
        rr = 5 + len(kp)
        va.cell(rr, 1, label)
        c = va.cell(rr, 2, "=" + formula)
        c.number_format = fmt
        c.font = BOLD
        va.cell(rr, 3, unit)
        va.cell(rr, 4, how).font = GREY
        kp[key] = f"Valuation!$B${rr}"
        wb.defined_names["kpi_" + key] = DefinedName("kpi_" + key, attr_text=f"Valuation!$B${rr}")
        return rr

    kpi("npv", "NPV at discount rate", f"SUMPRODUCT({RNG('fcf')},{RNG('df')})", "EUR", "sum of FCF x 1/(1+r)^t, t=0..10")
    kpi("npv_check", "NPV cross-check with Excel NPV()", f"{R('fcf', 0)}+NPV(discount_rate_real/100,{RNG('fcf', 1)})", "EUR", "t=0 + NPV(r, t1..t10)")
    kpi("discount_rate", "Discount rate (real)", "discount_rate_real/100", "%", "input", fmt=PCT)
    kpi("irr", "IRR (real)", f'IFERROR(IRR({RNG("fcf")},0.1),IFERROR(IRR({RNG("fcf")},-0.1),IFERROR(IRR({RNG("fcf")},-0.3),"n/a")))', "%", "IRR of FCF t=0..10 incl. salvage", fmt=PCT)
    kpi("payback", "Payback period (simple)", f'IF(COUNT({RNG("pb_h", 1)})=0,"not reached",MIN({RNG("pb_h", 1)}))', "years from t=0",
        "first year cumulative FCF (excl. salvage) >= 0, interpolated", fmt="0.0")
    kpi("dpayback", "Discounted payback period", f'IF(COUNT({RNG("dpb_h", 1)})=0,"not reached",MIN({RNG("dpb_h", 1)}))', "years from t=0",
        "same on discounted FCF", fmt="0.0")
    X3 = lambda key: f"INDEX({RNG(key)},1,opt_start_delay+4)"      # third operating year
    kpi("ebitda_y3", "EBITDA in 3rd operating year (2029; 2030 if sales start 2028)", X3("ebitda"), "EUR", "first year at full output (criterion 3)")
    kpi("ebitda_y1", "EBITDA year 1 (2027)", R("ebitda", 1), "EUR", "")
    kpi("ebitda_y10", "EBITDA year 10 (2036)", R("ebitda", 10), "EUR", "")
    kpi("funding", "Funding need (peak cumulative negative FCF)", f"MAX(0,-MIN({RNG('cum')}))", "EUR",
        "equity/debt needed before the plant pays for itself")
    kpi("capex", "Initial capex", f"SUM({RNG('capex_initial')})", "EUR", "plant + civil + permits (+hooks, crushing line)")
    kpi("npv_pct", "NPV as % of initial capex", f"{kp['npv']}/{kp['capex']}", "%", "borderline test: -10% .. 0%", fmt=PCT)
    kpi("npv_wc", "Memo: NPV if working capital is released in 2036", f"{kp['npv']}+{R('wc', 10)}*{R('df', 10)}", "EUR",
        "not used for the criteria (plan 2.1 allows salvage only)")
    kpi("rev_y3", "Revenue 3rd operating year", X3("revenue"), "EUR", "")
    kpi("unit_contrib", "Contribution per m3, 3rd operating year",
        f"({X3('revenue')}-{X3('rev_extra')}-{X3('rev_crusher')})/{X3('vol')}-({X3('opex_var')}-{X3('cr_var')})/{X3('vol')}",
        "EUR/m3", "concrete price - concrete variable costs (incl. bad debt, delivery)", fmt=EUR2)
    kpi("be_vol", "Break-even volume EBITDA = 0 (3rd-operating-year prices & fixed costs)",
        f'IF({kp["unit_contrib"]}>0,({X3("opex_fixed")}-{X3("rev_extra")}-{X3("crusher_net")})/{kp["unit_contrib"]},"none")', "m3/yr",
        "(fixed costs - hook revenue - crusher net) / contribution per m3")
    kpi("cap_charge", "Capital recovery charge (10 yr annuity of initial capex at r)",
        f"{kp['capex']}*(discount_rate_real/100)/(1-(1+discount_rate_real/100)^-{NYRS})", "EUR/yr", "memo")
    kpi("be_vol_cc", "Volume where EBITDA covers the capital charge (year-3 basis, pre-tax)",
        f'IF({kp["unit_contrib"]}>0,({X3("opex_fixed")}-{X3("rev_extra")}-{X3("crusher_net")}+{kp["cap_charge"]})/{kp["unit_contrib"]},"none")', "m3/yr",
        "rough NPV break-even; exact Python value below")
    va.column_dimensions["A"].width = 62
    va.column_dimensions["B"].width = 16
    va.column_dimensions["C"].width = 14
    va.column_dimensions["D"].width = 60
    va_py_row = 5 + len(kp) + 1

    # ======================= UnitEconomics ==================================
    ue = sheets["UnitEconomics"]
    ue["A1"] = "Unit economics - EUR per m3 delivered, circular vs standard (conventional) mix in the same plant"
    ue["A1"].font = TITLE
    ue["A3"] = "Year index shown (1..10; 3 = 2029, first full-output year)"
    ue["C3"] = 3
    ue["C3"].font, ue["C3"].fill = BLUE, INPUT_FILL
    ue["D3"] = f"={mc.BASE_YEAR}+C3"
    ue["A4"] = "Fixed costs are spread equally over all m3 of that year. Delivery is subcontracted."
    ue["A4"].font = GREY
    hdr(ue, 6, ["Line (EUR/m3)", "", "Circular", "Standard CEM III/B", "Plant average", "Circular - standard"])

    def ix(key):
        sh, row = REG[key]
        return f"INDEX({sh}!$D${row}:$M${row},1,$C$3)"

    share = f"({ix('vol_circ')}/{ix('vol')})"
    ue_rows = []
    def ue_row(label, circ, conv, avg=None, total=False):
        rr = 7 + len(ue_rows)
        ue.cell(rr, 1, label)
        ue.cell(rr, 3, "=" + circ)
        ue.cell(rr, 4, "=" + conv)
        ue.cell(rr, 5, "=" + (avg if avg else f"{share}*C{rr}+(1-{share})*D{rr}"))
        ue.cell(rr, 6, f"=C{rr}-D{rr}")
        for cc_ in range(3, 7):
            ue.cell(rr, cc_).number_format = EUR2
            if total:
                ue.cell(rr, cc_).font = BOLD
                ue.cell(rr, cc_).fill = TOTAL_FILL
        if total:
            ue.cell(rr, 1).font = BOLD
        ue_rows.append(rr)
        return rr

    rp = ue_row("Price (delivered)", ix("price_circ"), ix("price_conv"), f"({ix('rev_circ')}+{ix('rev_conv')})/{ix('vol')}", total=True)
    first = ue_row("Cement", ix("c_cement"), ix("v_cement"))
    ue_row("Gravel", ix("c_gravel"), ix("v_gravel"))
    ue_row("Recycled aggregate (delivered)", ix("c_rca"), ix("v_rca"))
    ue_row("Sand", ix("c_sand"), ix("v_sand"))
    ue_row("Admixture (superplasticiser)", ix("c_admix"), ix("v_admix"))
    ue_row("Water", ix("c_water"), ix("v_water"))
    ue_row("RCA testing & handling", f"Opex!$C${REG['rca_qc_m3'][1]}", "0")
    ue_row("Energy", f"Opex!$C${REG['energy_m3'][1]}", f"Opex!$C${REG['energy_m3'][1]}")
    ue_row("Variable cost delta (hook)", "opt_cost_delta_per_m3", "opt_cost_delta_per_m3")
    ue_row("Circular-only cost delta (option)", f"IF({ix('vol_circ')}>0,{ix('o_circdelta')}/{ix('vol_circ')},0)", "0")
    crn = f"-({ix('crusher_net')}-{ix('rev_crusher')}*bad_debt_pct_revenue/100)/{ix('vol')}"
    ue_row("On-site crushing net cost per m3 (negative = net income, option A+B)", crn, crn)
    fixed_keys = [("o_staffprod", "Production labour"), ("o_maint", "Maintenance & wear"),
                  ("o_cert", "Certification & testing"), ("o_lease", "Site lease"),
                  ("o_staffsga", "SG&A staff"), ("o_sgaother", "SG&A non-staff"), ("o_extrafixed", "Extra fixed (hook)")]
    for k, lab in fixed_keys:
        f_ = f"{ix(k)}/{ix('vol')}"
        ue_row(lab, f_, f_)
    ue_row("Delivery (truck mixer)", f"Opex!$C${REG['delivery_m3'][1]}", f"Opex!$C${REG['delivery_m3'][1]}")
    ue_row("Bad debt", f"C{rp}*bad_debt_pct_revenue/100", f"D{rp}*bad_debt_pct_revenue/100")
    last = ue_rows[-1]
    rc = ue_row("Cash cost (before depreciation)", f"SUM(C{first}:C{last})", f"SUM(D{first}:D{last})", f"SUM(E{first}:E{last})", total=True)
    re_ = ue_row("EBITDA margin", f"C{rp}-C{rc}", f"D{rp}-D{rc}", f"E{rp}-E{rc}", total=True)
    rd = ue_row("Depreciation & amortisation", f"{ix('da')}/{ix('vol')}", f"{ix('da')}/{ix('vol')}")
    rt = ue_row("Total cost incl. depreciation", f"C{rc}+C{rd}", f"D{rc}+D{rd}", f"E{rc}+E{rd}", total=True)
    rb = ue_row("EBIT margin", f"C{rp}-C{rt}", f"D{rp}-D{rt}", f"E{rp}-E{rt}", total=True)
    rk = ue_row("Memo: capital recovery charge (10 yr annuity at r)", f"{kp['cap_charge']}/{ix('vol')}", f"{kp['cap_charge']}/{ix('vol')}")
    ue_row("Memo: EBITDA margin after capital charge", f"C{re_}-C{rk}", f"D{re_}-D{rk}", f"E{re_}-E{rk}", total=True)
    rr = ue_rows[-1] + 2
    ue.cell(rr, 1, "Check: plant-average EBITDA x volume - P&L EBITDA (0 unless hook revenue)")
    ue.cell(rr, 5, f"=E{re_}*{ix('vol')}-({ix('ebitda')}-{ix('rev_extra')})").number_format = EUR2
    ue.cell(rr, 5).fill = CHECK_FILL
    ue.cell(rr + 1, 1, "EBITDA margin as % of price")
    for cc_, L in ((3, "C"), (4, "D"), (5, "E")):
        ue.cell(rr + 1, cc_, f"={L}{re_}/{L}{rp}").number_format = PCT
    ue.column_dimensions["A"].width = 52
    for L in "CDEF":
        ue.column_dimensions[L].width = 18

    # ======================= Python-computed sheets ========================
    tab = table
    res = {sc: mc.run(tab, scenario=sc) for sc in mc.SCENARIOS}
    sens = mc.sensitivity(tab)
    viab = mc.viability(tab)
    extra = python_extras(tab)

    # Valuation: Python-only extras
    va.cell(va_py_row, 1, "Computed by model_core.py (base case, values - not live)").font = Font(bold=True, color="7F7F7F")
    for i, (lab, val, unit) in enumerate(extra["rows"]):
        c1 = va.cell(va_py_row + 1 + i, 1, lab)
        c2 = va.cell(va_py_row + 1 + i, 2, val if val is not None else "not reached")
        va.cell(va_py_row + 1 + i, 3, unit)
        c1.font = c2.font = GREY
        c2.number_format = EUR if abs(val or 0) > 100 else EUR2

    ws = sheets["Scenarios"]
    ws["A1"] = "Scenarios - pessimistic / base / optimistic (values from model_core.py)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Each input takes the column shown in Inputs!F (pessimistic) and Inputs!G (optimistic). Cost-type inputs: "
                "pessimistic = high column; revenue-type inputs: pessimistic = low column; 'fixed' inputs (discount rate, tax, "
                "lifetimes, recipes) stay at base. ALL inputs move together, so these are stress bounds, not likely cases. "
                "To see a scenario live, set Inputs!C3 to 1, 2 or 3.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 48
    hdr(ws, 4, ["KPI", "Pessimistic", "Base", "Optimistic", "unit"])
    kpis = [("NPV @ 8%", "npv", EUR, "EUR"), ("IRR", "irr", PCT, "%"), ("Payback", "payback", "0.0", "yr"),
            ("Discounted payback", "discounted_payback", "0.0", "yr"), ("EBITDA year 3 (2029)", "ebitda_y3", EUR, "EUR"),
            ("Funding need", "funding_need", EUR, "EUR"), ("Initial capex", "capex_initial", EUR, "EUR"),
            ("NPV / capex", "npv_pct_of_capex", PCT, "%"), ("Break-even volume (EBITDA=0, yr-3 basis)", "be_volume_ebitda_y3", EUR, "m3/yr")]
    for i, (lab, k, fmt, unit) in enumerate(kpis):
        ws.cell(5 + i, 1, lab)
        for j, sc in enumerate(mc.SCENARIOS):
            v = res[sc]["kpis"][k]
            c = ws.cell(5 + i, 2 + j, v if v is not None else "n/a / not reached")
            c.number_format, c.font = fmt, GREY
        ws.cell(5 + i, 5, unit)
    r0 = 6 + len(kpis)
    for blk, key in (("EBITDA by year", "ebitda"), ("Free cash flow by year", "fcf"), ("Volume by year", "volume")):
        ws.cell(r0, 1, blk).font = BOLD
        for t in mc.T:
            ws.cell(r0, 2 + t, mc.YEARS[t]).font = BOLD
        for j, sc in enumerate(mc.SCENARIOS):
            ws.cell(r0 + 1 + j, 1, sc)
            for t in mc.T:
                c = ws.cell(r0 + 1 + j, 2 + t, res[sc]["years"][key][t])
                c.number_format, c.font = EUR, GREY
        r0 += 5
    ws.cell(r0, 1, "Plausible range (memo): base NPV +/- root-sum-square of the one-at-a-time NPV deviations of independent drivers").font = BOLD
    ws.cell(r0 + 1, 1, "NPV low (RSS)"); ws.cell(r0 + 1, 2, extra["rss_low"]).number_format = EUR
    ws.cell(r0 + 2, 1, "NPV high (RSS)"); ws.cell(r0 + 2, 2, extra["rss_high"]).number_format = EUR
    ws.cell(r0 + 3, 1, "Drivers in RSS: " + ", ".join(extra["rss_drivers"])).font = GREY
    ws.column_dimensions["A"].width = 44
    for c_ in range(2, 14):
        ws.column_dimensions[get_column_letter(c_)].width = 14

    ws = sheets["Sensitivity"]
    ws["A1"] = "Sensitivity - one input group at a time to its low / high column, effect on NPV @ 8% (values from model_core.py)"
    ws["A1"].font = TITLE
    ws["A2"] = f"Base NPV = {res['base']['kpis']['npv']:,.0f} EUR. Sorted by swing (tornado order)."
    hdr(ws, 4, ["Rank", "Driver", "NPV at LOW input", "NPV at HIGH input", "Swing", "Pessimistic column",
                "NPV at pessimistic", "Main driver (crit. 4)", "Parameters moved"])
    for i, s in enumerate(sorted(sens, key=lambda s: -s["swing"])):
        vals = [i + 1, s["name"], s["npv_low_input"], s["npv_high_input"], s["swing"], s["pessimistic_column"],
                s["npv_pessimistic"], "yes" if s["main_driver"] else "", ", ".join(s["params"][:3]) + (" ..." if len(s["params"]) > 3 else "")]
        for j, v in enumerate(vals):
            c = ws.cell(5 + i, 1 + j, v)
            c.font = GREY
            if j in (2, 3, 4, 6):
                c.number_format = EUR
    for c_, w in zip("ABCDEFGHI", (6, 56, 16, 16, 14, 12, 16, 12, 70)):
        ws.column_dimensions[c_].width = w

    ws = sheets["Viability"]
    ws["A1"] = "Viability gate - plan.md section 2.1 (criteria fixed before results)"
    ws["A1"].font = TITLE
    ws["A2"] = '="Criteria 1-3 are LIVE for the active scenario: "&Inputs!D3&". Criterion 4 is the base-case robustness test from model_core.py."'
    hdr(ws, 4, ["#", "Criterion", "Value", "Threshold", "Result"])
    ws["A5"], ws["B5"], ws["C5"], ws["D5"] = 1, "NPV > 0 at 8% real (10 yr, salvage-only terminal value)", f"={kp['npv']}", "> 0"
    ws["E5"] = f'=IF(C5>0,"PASS","FAIL")'
    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = 2, "Payback period < 7 years", f"={kp['payback']}", "< 7"
    ws["E6"] = '=IF(ISNUMBER(C6),IF(C6<7,"PASS","FAIL"),"FAIL")'
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = 3, "EBITDA positive by year 3 (2029)", f"={kp['ebitda_y3']}", "> 0"
    ws["E7"] = '=IF(C7>0,"PASS","FAIL")'
    ws["A8"], ws["B8"] = 4, "NPV > 0 with each of 6 main drivers at its pessimistic value (one at a time): at least 2 of 6"
    n4 = viab["c4_n_pass"]
    ws["C8"], ws["D8"] = n4, ">= 2 of 6"
    ws["C8"].font = GREY
    ws["E8"] = '=IF(C8>=2,"PASS","FAIL")'
    ws["A10"], ws["B10"] = "", "Borderline test: NPV between -10% and 0% of initial capex"
    ws["C10"] = f"={kp['npv_pct']}"
    ws["C10"].number_format = PCT
    ws["E10"] = '=IF(AND(C10>=-0.1,C10<=0),"YES","NO")'
    ws["B12"] = "VERDICT"
    ws["B12"].font = Font(bold=True, size=13)
    ws["C12"] = ('=IF(AND(E5="PASS",E6="PASS",E7="PASS",E8="PASS"),"VIABLE",'
                 'IF(OR(AND(E5="PASS",E6="PASS",E7="PASS"),E10="YES"),"BORDERLINE","NOT VIABLE"))')
    ws["C12"].font = Font(bold=True, size=13, color="C00000")
    for rr_ in range(5, 9):
        ws.cell(rr_, 3).number_format = EUR if rr_ != 6 else "0.0"
    hdr(ws, 14, ["", "Criterion 4 detail (base case, model_core.py)", "NPV at pessimistic", "Pessimistic column", "NPV > 0?"])
    for i, s in enumerate(viab["c4_detail"]):
        ws.cell(15 + i, 2, s["name"])
        c = ws.cell(15 + i, 3, s["npv_pessimistic"]); c.number_format = EUR; c.font = GREY
        ws.cell(15 + i, 4, s["pessimistic_column"])
        ws.cell(15 + i, 5, "PASS" if s["npv_pessimistic"] > 0 else "FAIL")
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 12

    # ======================= Opt_Overview (Agent 7) ===========================
    ws = sheets["Opt_Overview"]
    ws["A1"] = "Viability options - live option selector (Agent 7)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Set Inputs!I3 to an option number (1..N, row 5) and the WHOLE model recalculates live with that option's values "
                "(blank cell = parameter keeps its scenario/base value). 0 = base case. Option O6 (asset-light) is a separate model: "
                "sheet Opt_AssetLight. Results table below: values from model/options.py (grey) - see research/viability_options.md.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 45
    ws["A3"] = "Active option number (Inputs!I3)"
    ws["B3"] = "=option_sel"
    ws["C3"] = f'=IF(option_sel=0,"none (base case)",INDEX($C$6:${OPT_C1}$6,1,option_sel))'
    ws["A3"].font = ws["B3"].font = ws["C3"].font = BOLD
    ws["E3"] = "Live NPV of active option"
    ws["G3"] = f"={kp['npv']}"
    ws["G3"].number_format = EUR
    ws["G3"].font = BOLD
    hdr(ws, 5, ["parameter", "base value"] + [str(i + 1) for i in range(len(opts.MODEL_OPTIONS))])
    ws.cell(6, 1, "option").font = BOLD
    ws.cell(6, 2, "")
    for j, o in enumerate(opts.MODEL_OPTIONS):
        c = ws.cell(6, 3 + j, f"{o[0]} {o[1]}")
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[6].height = 75
    hook_names = set(mc.HOOKS)
    for i, k_ in enumerate(opt_params):
        rr = OPT_R0 + i
        ws.cell(rr, 1, k_)
        ws.cell(rr, 2, 0.0 if k_ in hook_names else table[k_]["base"]).font = GREY
        for j, o in enumerate(opts.MODEL_OPTIONS):
            if k_ in o[3]:
                c = ws.cell(rr, 3 + j, o[3][k_])
                c.font, c.fill = BLUE, INPUT_FILL
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12
    for j in range(len(opts.MODEL_OPTIONS)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 15
    if with_python_sheets:
        ores = opts.evaluate(table)
        r0 = OPT_R1 + 3
        ws.cell(r0, 1, "Results per option (values from model/options.py; criterion 4 = option-relative stress test)").font = BOLD
        hdr(ws, r0 + 1, ["option", "NPV @ 8%", "IRR", "payback (yr)", "EBITDA 3rd op. year", "funding need", "initial capex",
                         "criteria passed", "crit. 4 (of 6)", "verdict", "feasibility 1-5"])
        for i, (key, rec) in enumerate(ores["options"].items()):
            k = rec["kpis"]
            vals = [f"{key} {rec['name']}", k["npv"], k["irr"] if k["irr"] is not None else "n/a",
                    k["payback"] if k["payback"] is not None else "not reached", k["ebitda_y3"], k["funding_need"],
                    k["capex_initial"], f"{rec['criteria']['n_pass']}/4", rec["criteria"]["c4_n_pass"], rec["criteria"]["verdict"],
                    rec.get("feasibility", "")]
            for j, v in enumerate(vals):
                c = ws.cell(r0 + 2 + i, 1 + j, v)
                c.font = GREY
                c.number_format = PCT if j == 2 else ("0.0" if j == 3 else EUR)
        r1 = r0 + 3 + len(ores["options"])
        ws.cell(r1, 1, "Policy dependence (values from model/options.py)").font = BOLD
        hdr(ws, r1 + 1, ["case", "levy on gravel+sand for NPV=0 (EUR/t)", "public green premium for NPV=0 (%)",
                         "MKI value on circular m3 for NPV=0 (EUR/m3)", "NPV high ETS path", "NPV gate fee 15", "price for NPV=0 (EUR/m3)",
                         "volume for NPV=0 (m3/yr)"])
        for i, pt in enumerate(ores["policy"]):
            vals = [pt["option"], pt["levy_for_npv0_eur_per_t"], pt["public_premium_for_npv0_pct"], pt["circ_mki_value_for_npv0_eur_m3"],
                    pt["npv_high_ets_path"], pt["npv_gate_fee_15"], pt["price_for_npv0_eur_m3"], pt["volume_for_npv0_m3"]]
            for j, v in enumerate(vals):
                c = ws.cell(r1 + 2 + i, 1 + j, "none in range" if v is None else v)
                c.font = GREY
                c.number_format = EUR2 if j in (1, 2, 3, 6) else EUR

    # ======================= Opt_AssetLight (O6, live) ========================
    al = YearSheet(sheets["Opt_AssetLight"], "Option O6 - asset-light licence + RCA service (archetype D), EUR real 2026",
                   "Separate small model (no plant). Inputs in column P (blue). Same tax and discounting rules as the plant model.")
    wsa = sheets["Opt_AssetLight"]
    hdr(wsa, 5, ["input", "", "value", "unit", "low", "high", "ID"], c0=15)
    for i, (k_, lab, unit, lo, b, hi, sid) in enumerate(opts.AL_INPUTS):
        rr = 6 + i
        wsa.cell(rr, 15, lab)
        c = wsa.cell(rr, 17, b)
        c.font, c.fill = BLUE, INPUT_FILL
        wsa.cell(rr, 18, unit)
        wsa.cell(rr, 19, lo)
        wsa.cell(rr, 20, hi)
        wsa.cell(rr, 21, sid)
        wb.defined_names[k_] = DefinedName(k_, attr_text=f"Opt_AssetLight!$Q${rr}")
    wsa.column_dimensions["O"].width = 58
    al.row("al_lic", "Licensee plants", "plants", "IF({t}=0,0,MIN(al_licensees_max,{t}))", fmt="0")
    al.row("al_vol", "Circular volume at licensees", "m3", "{al_lic}*al_circ_volume_per_licensee")
    al.row("al_rfee", "Licence / service fees", "EUR", "{al_vol}*al_fee_per_m3")
    al.row("al_rrca", "RCA brokerage margin", "EUR", "{al_vol}*al_rca_t_per_m3*al_rca_margin_per_t")
    al.row("al_rsrv", "LCA/EPD/MKI services", "EUR", "IF({t}>=2,al_services_rev,0)")
    al.row("al_rev", "Revenue", "EUR", "{al_rfee}+{al_rrca}+{al_rsrv}", total=True)
    al.row("al_staff", "Staff", "EUR", "IF({t}>=1,al_fte*al_cost_per_fte,0)")
    al.row("al_ovh", "Other overhead", "EUR", "IF({t}>=1,al_overhead,0)")
    al.row("al_start", "Start-up costs (expensed)", "EUR", "IF({t}=0,al_startup_t0,0)")
    al.row("al_ebitda", "EBITDA", "EUR", "{al_rev}-{al_staff}-{al_ovh}-{al_start}", total=True)
    al.row("al_dep", "Depreciation", "EUR", "IF(AND({t}>=1,{t}<=al_capex_life),al_capex_t0/al_capex_life,0)")
    al.row("al_ebit", "EBIT", "EUR", "{al_ebitda}-{al_dep}", total=True)
    al.row("al_taxable", "Taxable profit after loss offset", "EUR",
           lambda t: "MAX(0,{al_ebit})" if t == 0 else "MAX(0,{al_ebit}-{al_pool@-1})")
    al.row("al_pool", "Tax-loss pool (end of year)", "EUR", lambda t: "MAX(0,-{al_ebit})" if t == 0 else "MAX(0,{al_pool@-1}-{al_ebit})")
    al.row("al_tax", "Corporate income tax", "EUR",
           "MIN({al_taxable},tax_vpb_low_bracket_limit)*tax_rate_vpb_low_bracket/100+MAX(0,{al_taxable}-tax_vpb_low_bracket_limit)*tax_rate_vpb_high_bracket/100")
    al.row("al_wc", "Receivables", "EUR", "{al_rev}*al_receivable_days/365")
    al.row("al_dwc", "Change in working capital (cash out = negative)", "EUR", lambda t: "-{al_wc}" if t == 0 else "-({al_wc}-{al_wc@-1})")
    al.row("al_capex", "Capex", "EUR", "IF({t}=0,-al_capex_t0,0)")
    al.row("al_fcf", "FREE CASH FLOW", "EUR", "{al_ebitda}-{al_tax}+{al_dwc}+{al_capex}", total=True)
    al.row("al_cum", "Cumulative FCF", "EUR", "SUM({al_fcf:cum})")
    al.row("al_df", "Discount factor", "x", "1/(1+discount_rate_real/100)^{t}", fmt="0.0000")
    al.row("al_pb", "Payback crossing (year)", "yr",
           lambda t: None if t == 0 else 'IF(AND({al_cum@-1}<0,{al_cum}>=0),{t}-1+(-{al_cum@-1})/({al_cum}-{al_cum@-1}),"")', fmt="0.00")
    al.write()
    rr = al.r + 1
    def alkpi(key, label, formula, fmt=EUR):
        global_r = rr + len([n for n in wb.defined_names if str(n).startswith("kpi_al_")])
        wsa.cell(global_r, 1, label).font = BOLD
        c = wsa.cell(global_r, 3, "=" + formula)
        c.number_format, c.font = fmt, BOLD
        wb.defined_names["kpi_al_" + key] = DefinedName("kpi_al_" + key, attr_text=f"Opt_AssetLight!$C${global_r}")
    alkpi("npv", "NPV @ discount rate", f"SUMPRODUCT({RNG('al_fcf')},{RNG('al_df')})")
    alkpi("irr", "IRR", f'IFERROR(IRR({RNG("al_fcf")},0.1),IFERROR(IRR({RNG("al_fcf")},-0.1),"n/a"))', PCT)
    alkpi("payback", "Payback (years from 2026)", f'IF(COUNT({RNG("al_pb", 1)})=0,"not reached",MIN({RNG("al_pb", 1)}))', "0.0")
    alkpi("ebitda_y3", "EBITDA 2029", R("al_ebitda", 3))
    alkpi("funding", "Funding need", f"MAX(0,-MIN({RNG('al_cum')}))")

    write_readme(sheets["README"])
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def python_extras(tab):
    base = mc.run(tab)
    k = base["kpis"]
    be_vol = mc.solve(tab, "output_steady_state", lo=20000, hi=240000)
    be_price = mc.solve(tab, "price_plant_avg_all_classes", lo=80, hi=250)
    be_prem = mc.solve(tab, "green_premium_public", lo=-50, hi=200, overrides={"green_premium_private": 0})
    be_capex = mc.solve(tab, "capex_plant_equipment", lo=-20, hi=20)
    rows = [
        ("NPV break-even steady-state volume (same ramp %)", be_vol, "m3/yr"),
        ("NPV break-even plant-average price (2026 basis)", be_price, "EUR/m3"),
        ("NPV break-even public green premium (private 0%)", be_prem, "% of price"),
        ("NPV break-even plant-equipment capex", be_capex, "MEUR (negative = impossible)"),
        ("NPV with 60k m3/yr steady state (market SOM)", mc.run(tab, {"output_steady_state": 60000})["kpis"]["npv"], "EUR"),
        ("NPV at 6% discount rate", mc.run(tab, {"discount_rate_real": 6})["kpis"]["npv"], "EUR"),
        ("NPV at 10% discount rate", mc.run(tab, {"discount_rate_real": 10})["kpis"]["npv"], "EUR"),
        ("NPV without grant", mc.run(tab, {"grant_dei_veki_first_plant": 0})["kpis"]["npv"], "EUR"),
    ]
    sens = mc.sensitivity(tab)
    indep = {"Green premium (public & private)", "Average selling price (EUR 115-155/m3)",
             "Recycled aggregate price, ex-works (EUR 10-25/t)", "Concrete delivery distance (10-35 km)",
             "RCA haul distance (15-50 km)", "Cement price 2026 (EUR 95-130/t)", "ETS/CBAM carbon cost path only",
             "Plant utilisation (50k-90k m3/yr)", "Ramp-up speed (yr1-3)", "Capex (plant, civil, permits)",
             "Gravel price (EUR 25-40/t)", "Sand price (EUR 14-25/t)", "Delivery cost rates (fixed + per km)",
             "Production staff (FTE x cost)", "SG&A (staff + non-staff)", "Site lease (area x rate)",
             "Circular share of volume (40-80%)", "Grant DEI+/VEKI (EUR 0-1.5M)", "Working capital days",
             "Salvage value (10-40% of capex)", "Maintenance (1.5-4% of capex)"}
    dn, up = 0.0, 0.0
    for s in sens:
        if s["name"] in indep:
            lo_ = min(s["npv_low_input"], s["npv_high_input"]) - s["npv_base"]
            hi_ = max(s["npv_low_input"], s["npv_high_input"]) - s["npv_base"]
            dn += min(0, lo_) ** 2
            up += max(0, hi_) ** 2
    return {"rows": rows, "rss_low": k["npv"] - math.sqrt(dn), "rss_high": k["npv"] + math.sqrt(up),
            "rss_drivers": sorted(indep), "be_vol": be_vol, "be_price": be_price, "be_prem": be_prem, "be_capex": be_capex}


def write_readme(ws):
    lines = [
        ("Financial model - B1 circular ready-mix plant, Zuid-Holland (Rotterdam-Drechtsteden)", TITLE),
        ("Built by model/build_model.py from model/inputs.csv. Same model as model/model_core.py (checked: model/check_report.md).", None),
        ("", None),
        ("COLOUR CODE", BOLD),
        ("Blue font (yellow fill) = input you may change. Black font = live formula. Grey italic = value computed by model_core.py (not live). Green fill = check (must be 0).", None),
        ("", None),
        ("HOW IT WORKS", BOLD),
        ("1. Inputs: every parameter has low/base/high. Cell Inputs!C3 picks the scenario (1 pessimistic, 2 base, 3 optimistic).", None),
        ("   Column H picks the column per parameter; column I is an OVERRIDE; column J is the ACTIVE value. Each active cell has a defined name = parameter name.", None),
        ("   Pessimistic = high column for cost-type inputs, low column for revenue-type inputs (column N 'direction'). 'fixed' inputs stay at base in all scenarios.", None),
        ("2. Capex -> Revenue -> Opex -> PnL -> CashFlow -> Valuation. Yearly sheets: column C = 2026 (t=0), D..M = 2027..2036 (t=1..10).", None),
        ("3. UnitEconomics: EUR/m3 stack for circular vs standard mix and plant average, for the year in UnitEconomics!C3.", None),
        ("4. Scenarios, Sensitivity, Viability criterion 4: computed with model_core.py (one model run per case) and pasted as values.", None),
        ("", None),
        ("KEY CONVENTIONS", BOLD),
        ("Real 2026 EUR excl. VAT, no general inflation. Discount rate 8% real (6% / 10% tested in Sensitivity).", None),
        ("Timing: 2026 (t=0) = permitting + construction; all initial capex, start-up costs (pre-opening, LCA/EPD) and grant in t=0, undiscounted.", None),
        ("   Operations 2027-2036 (t=1..10). Ramp: 50% / 79% / 100% of 70,000 m3 in 2027 / 2028 / 2029+. NOTE: starting sales in 2027 is optimistic", None),
        ("   (technology.md: 12-18 months build incl. permit from late 2026). A one-year delay is not in the base case.", None),
        ("Price/cost basis: plant average over all strength classes (EUR 135/m3, ~310 kg cement), NOT C30/37 (Agent 5 advice).", None),
        ("Concrete price follows 80% of the market-wide CEM III/B carbon & slag cost increase. Green premium only on circular volume.", None),
        ("Tax: Vpb 19% to EUR 200k, 25.8% above; losses carried forward; MIA as extra deduction in t=0; grant taxable in t=0. Unlevered (no debt).", None),
        ("Terminal value = salvage value of plant+civil only (25%), tax-neutral, in 2036. Working capital NOT released (memo on Valuation).", None),
        ("Payback uses cumulative FCF excluding salvage, interpolated; measured from t=0 (2026).", None),
        ("Funding need = peak cumulative negative free cash flow (before any financing).", None),
        ("", None),
        ("ADDING OPTION SCENARIOS (for Agent 7)", BOLD),
        ("a) Fill the Override column (Inputs!I) or the option hooks (opt_*, bottom of Inputs). Nothing else needs to change.", None),
        ("b) In Python: model_core.run(load_inputs(), overrides={...}); build a copy of this workbook with build_model.build_workbook(path, overrides={...}).", None),
        ("c) For a new option sheet: add a sheet that reads the defined names / kpi_* names (e.g. kpi_npv) or paste model_core results.", None),
        ("VIABILITY OPTIONS (Agent 7): Inputs!I3 = option number (Opt_Overview lists them); the whole model recalculates live with that option.", BOLD),
        ("   New live logic (0 = off, base unchanged): opt_start_delay (first sales later), on-site crushing block (opt_crusher_*, hybrid_* inputs),", None),
        ("   circular-only price/cost deltas from a start year. Option O6 asset-light = separate live model on sheet Opt_AssetLight.", None),
        ("", None),
        ("FILES", BOLD),
        ("model/inputs.csv, model/model_core.py, model/build_model.py, model/check_model.py, model/charts.py, model/charts/*.png, model/model_summary.md, model/check_report.md", None),
    ]
    for i, (txt, font) in enumerate(lines):
        c = ws.cell(1 + i, 1, txt)
        if font:
            c.font = font
    ws["A6"].font = BLUE
    ws.column_dimensions["A"].width = 160


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--scenario", type=int, default=2)
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--no-recalc", action="store_true", help="skip LibreOffice recalculation (no cached values)")
    a = ap.parse_args()
    path = build_workbook(a.out, a.scenario)
    print("written", path)
    if not a.no_recalc and shutil.which("soffice"):
        # Recalculate once with LibreOffice so the file also carries cached values
        # (formulas are kept). Viewers that do not calculate then still show numbers.
        tmpd = tempfile.mkdtemp(prefix="cc_build_")
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmpd, path],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        shutil.copyfile(os.path.join(tmpd, os.path.basename(path)), path)
        shutil.rmtree(tmpd, ignore_errors=True)
        print("recalculated with LibreOffice (formulas kept, cached values added)")
